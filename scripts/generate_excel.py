#!/usr/bin/env python3
"""Export a normalized marketplace listing manifest to a reviewable workbook."""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from validate_listing import RULES, production_readiness, validate


HEADER_FILL = PatternFill("solid", fgColor="263A5B")
SUB_FILL = PatternFill("solid", fgColor="EAF0F8")
WRAP_TOP = Alignment(vertical="top", wrap_text=True)


def display(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if value is None:
        return ""
    return str(value)


def add_sheet(wb, name, title, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = HEADER_FILL
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="45658F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28
    for row in rows:
        ws.append([display(item) for item in row])
        row_index = ws.max_row
        for cell in ws[row_index]:
            cell.alignment = WRAP_TOP
            cell.font = Font(name="Arial", size=9, color="222222")
            if row_index % 2:
                cell.fill = SUB_FILL
        ws.row_dimensions[row_index].height = 34
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"
    return ws


def text_value(item):
    return item.get("text", "") if isinstance(item, dict) else (item or "")


def fact_references(item):
    return ", ".join(item.get("fact_ids", [])) if isinstance(item, dict) else ""


def build_workbook(manifest):
    platform_key = manifest.get("platform", "unknown")
    market = manifest.get("market", "unknown")
    profile = RULES.get("platforms", {}).get(platform_key, {})
    platform = profile.get("name", platform_key)
    product = manifest.get("product", {})
    category = manifest.get("category", {})
    listing = manifest.get("listing", {})
    wb = Workbook()
    wb.remove(wb.active)

    listing_rows = [
        ("Platform", platform), ("Market / site", market), ("Locale", manifest.get("locale")),
        ("Product", product.get("name")), ("Brand", product.get("brand")),
        ("Model", product.get("model")), ("Seller SKU", product.get("seller_sku")),
        ("GTIN", product.get("gtin")), ("Condition", product.get("condition")),
        ("Category", category.get("name")), ("Category ID", category.get("id")),
        ("Category schema", category.get("schema_source")),
        ("Title", text_value(listing.get("title"))),
        ("Title fact IDs", fact_references(listing.get("title"))),
        ("Description", text_value(listing.get("description"))),
        ("Description fact IDs", fact_references(listing.get("description"))),
    ]
    for index, item in enumerate(listing.get("highlights", []), start=1):
        listing_rows.extend([(f"Highlight {index}", text_value(item)), (f"Highlight {index} fact IDs", fact_references(item))])
    for index, item in enumerate(listing.get("disclosures", []), start=1):
        listing_rows.extend([(f"Disclosure {index}", text_value(item)), (f"Disclosure {index} fact IDs", fact_references(item))])
    for index, item in enumerate(listing.get("attributes", []), start=1):
        listing_rows.append((f"Attribute {index}", {key: value for key, value in item.items()}))
    if listing.get("search_terms"):
        listing_rows.append(("Search terms", listing.get("search_terms")))
    if listing.get("platform_fields"):
        listing_rows.append(("Platform fields", listing.get("platform_fields")))
    if product.get("variants"):
        listing_rows.append(("Variants", product.get("variants")))
    add_sheet(wb, "Listing", f"{platform} listing · {product.get('name', '')}", ["Field", "Value"], listing_rows, [28, 105])

    media_rows = []
    for item in manifest.get("media", []):
        media_rows.append((
            item.get("id"), item.get("role"), item.get("slot_name"), item.get("source_type"),
            item.get("path") or item.get("url"), item.get("width"), item.get("height"),
            item.get("bytes"), item.get("format"), item.get("background"),
            item.get("contains_text"), item.get("prompt"), ", ".join(item.get("fact_ids", [])),
            ", ".join(item.get("variant_ids", [])), item.get("buyer_question"),
            item.get("visual_proof"), item.get("layout"), item.get("on_image_text"),
            item.get("why_gallery_slot"),
        ))
    add_sheet(wb, "Media", "Image / video assets · specs and prompts", [
        "ID", "Role", "Slot", "Source type", "File / URL", "Width", "Height", "Bytes", "Format", "Background", "Text overlay", "Prompt / brief", "Fact IDs", "Variant IDs", "Buyer question", "Visual proof", "Layout", "Exact on-image text", "Why this slot"
    ], media_rows, [16, 15, 24, 16, 42, 10, 10, 14, 12, 18, 14, 70, 18, 18, 36, 42, 32, 35, 40])

    fact_rows = []
    for fact in manifest.get("facts", []):
        fact_rows.append((fact.get("id"), fact.get("field"), fact.get("value"), fact.get("source_type"), fact.get("confidence"), fact.get("evidence"), fact.get("source_url")))
    add_sheet(wb, "Product Facts", "Evidence ledger · claims must map back to these facts", [
        "Fact ID", "Field", "Value", "Source type", "Confidence", "Evidence / note", "Source URL"
    ], fact_rows, [16, 22, 36, 22, 14, 48, 48])

    sheet = manifest.get("selling_point_sheet", {})
    point_rows = [("Buyer question", question) for question in sheet.get("buyer_questions", [])]
    for index, point in enumerate(sheet.get("points", []), start=1):
        point_rows.extend([
            (f"Point {index}", point.get("point")),
            (f"Point {index} · evidence tier / fact IDs", f"{point.get('evidence_tier')} / {', '.join(point.get('fact_ids', []))}"),
            (f"Point {index} · buyer concern / benefit", f"{point.get('buyer_concern')} / {point.get('buyer_benefit', '')}"),
            (f"Point {index} · visual proof / claim risk", f"{point.get('visual_proof')} / {point.get('claim_risk', '')}"),
        ])
    point_rows.extend(("Unknown / do not invent", value) for value in sheet.get("unknowns", []))
    point_rows.extend(("Excluded claim", value) for value in sheet.get("excluded_claims", []))
    add_sheet(wb, "Selling Points", "Buyer concerns · evidenced points · proof plan", ["Planning item", "Value"], point_rows, [38, 110])

    findings = validate(manifest)
    readiness = production_readiness(manifest, findings)
    reviewed_on = RULES.get("reviewed_on")
    check_rows = [("Profile reviewed on", reviewed_on), ("Market", market), ("Production readiness", readiness["status"]), ("Readiness note", readiness["note"]), ("Category schema source", category.get("schema_source")), ("Category schema fetched at", category.get("schema_fetched_at")), ("Seller Center check status", manifest.get("policy_snapshot", {}).get("manual_check_status", "not_checked")), ("Seller Center checked at", manifest.get("policy_snapshot", {}).get("seller_center_checked_at"))]
    for field, value in manifest.get("release_review", {}).items():
        check_rows.append((f"Release review · {field}", value))
    for index, finding in enumerate(findings, start=1):
        source = finding.get("source", {})
        check_rows.append((f"{finding['level'].upper()} · {finding['code']}", finding["message"] + (f"\n{source.get('url')}" if source.get("url") else "")))
    if not findings:
        check_rows.append(("Preflight", "No automated findings. Complete the live Seller Center/category checks before publishing."))
    add_sheet(wb, "Preflight & Sources", "Policy preflight · warnings do not equal approval", ["Check", "Result / source"], check_rows, [34, 110])
    wb.properties.title = f"{platform} listing package"
    wb.properties.subject = f"Market {market}; rule snapshot {reviewed_on}"
    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("manifest", help="Path to listing manifest JSON")
    parser.add_argument("--output", help="Workbook path; defaults to platform_market_listing.xlsx")
    args = parser.parse_args()
    try:
        with Path(args.manifest).open(encoding="utf-8") as stream:
            manifest = json.load(stream)
    except (OSError, json.JSONDecodeError) as exc:
        print(f"Cannot read manifest: {exc}", file=sys.stderr)
        return 2
    if not isinstance(manifest, dict):
        print("Manifest root must be a JSON object.", file=sys.stderr)
        return 2
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{manifest.get('platform', 'listing')}_{manifest.get('market', 'market')}").strip("_")
    output = Path(args.output or f"{safe_name}_listing.xlsx")
    output.parent.mkdir(parents=True, exist_ok=True) if output.parent != Path(".") else None
    build_workbook(manifest).save(output)
    print(f"Workbook saved: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
